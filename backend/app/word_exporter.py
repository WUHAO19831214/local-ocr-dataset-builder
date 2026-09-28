"""Export OCR datasets to Word in visual and editable forms."""

from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt


class WordExportError(RuntimeError):
    pass


def export_word_files(pdf_path: Path, dataset_dir: Path, output_name: str) -> list[Path]:
    """Create both files beside the existing Markdown and JSON outputs."""
    pdf_path = pdf_path.expanduser().resolve()
    dataset_dir = dataset_dir.expanduser().resolve()
    markdown_path = dataset_dir / f"{output_name}.md"
    if not pdf_path.is_file() or pdf_path.suffix.lower() != ".pdf":
        raise WordExportError(f"找不到原始 PDF：{pdf_path}")
    if not markdown_path.is_file():
        raise WordExportError(f"找不到 Markdown：{markdown_path}")

    visual_path = dataset_dir / f"{output_name}_原版式.docx"
    editable_path = dataset_dir / f"{output_name}_可编辑.docx"
    export_visual_docx(pdf_path, visual_path)
    export_editable_docx(markdown_path, editable_path)
    return [visual_path, editable_path]


def export_visual_docx(pdf_path: Path, output_path: Path, dpi: int = 180) -> Path:
    """Place a rendered source PDF page on each Word page."""
    try:
        import pymupdf
    except ImportError as exc:
        raise WordExportError("缺少 PyMuPDF，无法生成原版式 Word。") from exc

    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Emu(0)
    section.left_margin = section.right_margin = Emu(0)
    section.header_distance = section.footer_distance = Emu(0)

    with pymupdf.open(str(pdf_path)) as source, tempfile.TemporaryDirectory() as temp_dir:
        if source.page_count == 0:
            raise WordExportError("PDF 没有页面。")
        first = source[0].rect
        if any(abs(page.rect.width - first.width) > 0.5 or abs(page.rect.height - first.height) > 0.5 for page in source):
            raise WordExportError("原版式 Word 暂不支持混合页面尺寸的 PDF。")
        section.page_width = Pt(first.width)
        section.page_height = Pt(first.height)
        for index, page in enumerate(source):
            image_path = Path(temp_dir) / f"page-{index + 1:04d}.png"
            page.get_pixmap(matrix=pymupdf.Matrix(dpi / 72, dpi / 72), alpha=False).save(str(image_path))
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(1)
            run = paragraph.add_run()
            run.font.size = Pt(1)
            shape = run.add_picture(str(image_path), width=Pt(first.width), height=Pt(first.height))
            _anchor_picture_to_page(shape)
            if index < source.page_count - 1:
                run.add_break(WD_BREAK.PAGE)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def _anchor_picture_to_page(shape) -> None:
    """Use absolute page positioning so Word cannot crop a full-page image."""
    inline = shape._inline
    anchor = OxmlElement("wp:anchor")
    for key, value in {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0",
        "simplePos": "0", "relativeHeight": "0", "behindDoc": "1",
        "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
    }.items():
        anchor.set(key, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for direction in ("H", "V"):
        position = OxmlElement(f"wp:position{direction}")
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = "0"
        position.append(offset)
        anchor.append(position)
    anchor.append(deepcopy(inline.extent))
    anchor.append(OxmlElement("wp:wrapNone"))
    anchor.append(deepcopy(inline.docPr))
    for frame_properties in inline.xpath("./wp:cNvGraphicFramePr"):
        anchor.append(deepcopy(frame_properties))
    anchor.append(deepcopy(inline.graphic))
    inline.getparent().replace(inline, anchor)


def export_editable_docx(markdown_path: Path, output_path: Path) -> Path:
    """Convert the recognized text, images and LaTeX math to editable Word."""
    pandoc = shutil.which("pandoc")
    if not pandoc:
        pandoc = next((str(path) for path in (Path("/usr/local/bin/pandoc"), Path("/opt/homebrew/bin/pandoc")) if path.is_file()), None)
    if pandoc:
        command = [
            pandoc,
            str(markdown_path),
            "--from=markdown+tex_math_dollars",
            "--to=docx",
            f"--resource-path={markdown_path.parent}",
            "--output",
            str(output_path),
        ]
        result = subprocess.run(command, capture_output=True, text=True, timeout=120)
        if result.returncode == 0 and output_path.is_file():
            _set_cjk_fonts(output_path)
            return output_path
        raise WordExportError(f"Pandoc 导出可编辑 Word 失败：{result.stderr.strip()}")

    # The bundled desktop app can still produce an editable text document when
    # Pandoc is unavailable. Formula source remains visible as LaTeX in this case.
    document = Document()
    for line in markdown_path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        image = re.fullmatch(r"!\[[^]]*\]\(([^)]+)\)", stripped)
        if image:
            path = (markdown_path.parent / image.group(1)).resolve()
            if path.is_relative_to(markdown_path.parent.resolve()) and path.is_file():
                document.add_picture(str(path), width=Pt(420))
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            document.add_heading(heading.group(2), level=min(len(heading.group(1)), 4))
        elif stripped:
            document.add_paragraph(stripped)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def _set_cjk_fonts(docx_path: Path) -> None:
    document = Document(docx_path)
    for root in (document.element, document.styles.element):
        for fonts in root.xpath(".//w:rFonts"):
            fonts.set(qn("w:eastAsia"), "PingFang SC")
    document.save(docx_path)
